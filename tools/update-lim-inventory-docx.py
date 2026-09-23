from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT_PATH = ROOT / "data1" / "LIM_CELL_INVENTORY.docx"
OUTPUT_PATH = DOCUMENT_PATH


def paragraph_by_text(document, text):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text == text)


def paragraph_index_by_text(document, text):
    return next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == text)


def set_single(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)


def set_labelled(paragraph, label, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    label_run = paragraph.add_run(f"{label}  ")
    label_run.bold = True
    paragraph.add_run(text)


def replace_intro_text(document, old, new, end_heading="Eight face learning mesh"):
    for paragraph in document.paragraphs:
        if paragraph.text == end_heading:
            break
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


document = Document(DOCUMENT_PATH)

set_single(document.paragraphs[1], "Current four archetypes and eight face learning mesh")
set_single(
    document.paragraphs[2],
    "This editable inventory records the Learning Information Mesh content used by the NourishlandXR introductory demo. It begins with four optional learning archetypes. Vision now belongs inside Shape the Outcome, alongside the cells that turn direction into goals, decisions and feedback. The inventory then lists the eight faces of the full mesh and their cells. Stable IDs, parent relationships and current companion text support safe content review and updates.",
)
set_single(
    document.paragraphs[3],
    "Source  app/services/limLearning.js    Verified against app/screens/temporaryArDemo.js    Updated 23 September 2026",
)
set_single(
    document.paragraphs[5],
    "Four archetypes form the introductory entry points: Read Nature, Understand the Land, Design the Forest and Shape the Outcome. Their cells remain optional during the demo and unfold only when an archetype is selected. Vision is a child of Shape the Outcome rather than a separate gate. The full LIM has 97 cells arranged under eight public faces. Ninety three retain the original Climate, Food forest, Plant and Pin IDs, while four face cells were added.",
)
set_single(document.paragraphs[8], "Four archetypes and introductory cells")

replace_intro_text(document, "Read the Place", "Read Nature")
replace_intro_text(document, "Read the place", "Read Nature")
replace_intro_text(document, "Understand Life", "Understand the Land")
replace_intro_text(document, "Understand life", "Understand the Land")

for stable_id in ["lim-intro-analysis", "lim-intro-literacy", "lim-intro-food-forest", "lim-intro-smart"]:
    paragraph = next(p for p in document.paragraphs if p.text.startswith(f"Stable ID  {stable_id}    "))
    set_single(paragraph, f"Stable ID  {stable_id}    Parent ID  none")

shape_heading = paragraph_by_text(document, "Shape the Outcome")
shape_index = paragraph_index_by_text(document, "Shape the Outcome")
shape_paragraphs = document.paragraphs
set_single(
    shape_paragraphs[shape_index + 2],
    "Shape the Outcome turns observations into a shared direction and useful decisions. Vision, clear goals, measurable outcomes, honest limits, practical choices and feedback keep technology connected to the living place it serves; the interface should inform care rather than replace it.",
)
set_labelled(shape_paragraphs[shape_index + 3], "Explore", "Vision · Goals · Outcomes · Limitations · Challenges · Decisions · Feedback")
set_labelled(shape_paragraphs[shape_index + 4], "Look for", "A future worth working toward, a decision that needs evidence, the people affected, a responsible person and the moment when the result should be reviewed.")
set_labelled(shape_paragraphs[shape_index + 5], "Ask", "What future should this work help create, what can the system help people notice, and what must remain a human judgement made in the place?")
set_labelled(shape_paragraphs[shape_index + 6], "Next", "Begin with Vision, connect it to a goal or decision, make uncertainty visible, then return through Feedback to Read Nature again.")

vision_heading = paragraph_by_text(document, "Vision")
vision_start = paragraph_index_by_text(document, "Vision")
read_heading = paragraph_by_text(document, "Read Nature")
vision_block = document.paragraphs[vision_start:paragraph_index_by_text(document, "Read Nature")]
vision_heading.style = document.styles["Heading 3"]
set_single(vision_block[1], "Stable ID  lim-intro-vision    Parent ID  lim-intro-smart")
set_single(
    vision_block[2],
    "Vision describes the future a project hopes to help create for a living place. It gives goals and decisions a shared direction without pretending the future is fixed; observation, participation and feedback can refine the vision as the place changes.",
)
set_labelled(vision_block[3], "Explore", "Purpose · Future state · Values · People · Place · Time horizon")
set_labelled(vision_block[4], "Look for", "The people and living systems included, the values guiding the work, a meaningful time horizon and signs that the imagined future still belongs to this place.")
set_labelled(vision_block[5], "Ask", "What future is worth working toward here, who should help shape it, and what would show that the vision needs to change?")
set_labelled(vision_block[6], "Next", "Connect Vision with Goals and Outcomes, then use Feedback to keep the shared direction responsive to the living place.")
set_labelled(vision_block[7], "Connect", "Goals · Outcomes")

goals_heading = paragraph_by_text(document, "Goals")
for paragraph in vision_block:
    goals_heading._p.addprevious(paragraph._p)

summary = document.tables[0]
summary._tbl.remove(summary.rows[1]._tr)
summary.rows[1].cells[0].text = "Four archetypes"
summary.rows[1].cells[1].text = "4"
summary.rows[1].cells[2].text = "Read Nature, Understand the Land, Design the Forest and Shape the Outcome"
summary.rows[2].cells[0].text = "Archetype topics"
summary.rows[2].cells[1].text = "23"
summary.rows[2].cells[2].text = "Detailed cells under the four archetypes, including Vision under Shape the Outcome"

paragraph_text = [paragraph.text for paragraph in document.paragraphs]
assert paragraph_text.index("Shape the Outcome") < paragraph_text.index("Vision") < paragraph_text.index("Goals")
assert not any("Parent ID  lim-intro-vision" in text for text in paragraph_text)
assert summary.rows[1].cells[0].text == "Four archetypes"
assert summary.rows[2].cells[1].text == "23"

document.save(OUTPUT_PATH)
print(OUTPUT_PATH)
